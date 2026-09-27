from __future__ import annotations

import copy
import io
import json
import re
import unicodedata
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any, Callable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Twips
from docx.table import Table
from docx.text.paragraph import Paragraph

SITE_URL = "https://www.simplexsciences.com"
CATALOG_PAGE_URL = f"{SITE_URL}/ssdna-ladders"
FEDEX_RATES_URL = "https://www.fedex.com/en-us/online/rating.html"

DOMESTIC_SERVICE = "FedEx 2-Day Shipping"
INTERNATIONAL_SERVICE = "FedEx International Priority Shipping"
SHIPPING_SERVICES = (
    "FedEx 2-Day Shipping",
    "FedEx Priority Overnight Shipping",
    "FedEx Standard Overnight Shipping",
    "FedEx Ground Shipping",
    "FedEx International Priority Shipping",
    "FedEx International Economy Shipping",
)

CENT = Decimal("0.01")
_PLAIN_OPTIONS = {"", "none", "standard"}


class QuoteTemplateError(ValueError):
    pass


def parse_money(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        amount = value
    elif isinstance(value, (int, float)):
        amount = Decimal(str(value))
    else:
        raw = str(value).strip().replace("$", "").replace(",", "").replace("USD", "").strip()
        if not raw:
            return None
        try:
            amount = Decimal(raw)
        except InvalidOperation:
            raise ValueError(f"{value!r} is not an amount") from None
    if not amount.is_finite() or amount < 0:
        raise ValueError(f"{value!r} is not an amount")
    return amount.quantize(CENT, rounding=ROUND_HALF_UP)


def format_money(amount: Decimal) -> str:
    return f"${amount:,.2f}"


@dataclass(frozen=True)
class Product:
    name: str
    url: str
    options: tuple[tuple[str, Decimal], ...]

    @property
    def option_labels(self) -> list[str]:
        return [label for label, _ in self.options]

    def price_for(self, option: str) -> Decimal:
        for label, price in self.options:
            if label == option:
                return price
        return self.options[0][1]

    def describe(self, option: str) -> str:
        if option.strip().lower() in _PLAIN_OPTIONS:
            return self.name
        return f"{self.name}, {option}"


BUNDLED_CATALOG_DATE = "2026-09-26"


def _p(name: str, slug: str, plain: str, dyed: str) -> Product:
    return Product(name, f"{SITE_URL}/products/p/{slug}", (("None", Decimal(plain)), ("Pre-Dyed", Decimal(dyed))))


BUNDLED_CATALOG: tuple[Product, ...] = (
    _p("ss10 DNA Ladder (100µL)", "ss10-dna-ladder", "85.00", "115.00"),
    _p("ss10 DNA Ladder 5-Pack (100µL x 5)", "ss10-dna-ladder-5-pack-2", "360.00", "510.00"),
    _p("ss20 DNA Ladder (100µL)", "ss20-dna-ladder-1", "169.00", "199.00"),
    _p("ss20 DNA Ladder 5-pack (100µL x 5)", "ss20-dna-ladder-2", "795.00", "795.00"),
    _p("ss50 DNA Ladder (100µL)", "ssdna-50", "459.00", "489.00"),
    _p("ss50+ DNA Ladder (100µL)", "ssdna-50-l3wsp", "689.00", "719.00"),
)


def _get_json(url: str, timeout: float) -> dict[str, Any]:
    request = urllib.request.Request(
        f"{url}?format=json",
        headers={"User-Agent": "Mozilla/5.0 (Simplex Sciences operations app)"},
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8"))


def _product_from_json(data: dict[str, Any], url: str) -> Product | None:
    item = data.get("item") or {}
    name = re.sub(r"\s+", " ", str(item.get("title") or "")).strip()
    options = []
    for variant in (item.get("structuredContent") or {}).get("variants") or []:
        money = variant.get("salePriceMoney") if variant.get("onSale") else variant.get("priceMoney")
        try:
            price = parse_money((money or {}).get("value"))
        except ValueError:
            price = None
        if price is None:
            continue
        label = " / ".join(str(v) for v in (variant.get("attributes") or {}).values()) or "Standard"
        options.append((label, price))
    if not name or not options:
        return None
    return Product(name, url, tuple(options))


def fetch_catalog(timeout: float = 4.0) -> tuple[Product, ...]:
    page = _get_json(CATALOG_PAGE_URL, timeout)
    paths = list(dict.fromkeys(re.findall(r'href="(/products/p/[^"?#]+)"', page.get("mainContent") or "")))
    if not paths:
        raise ValueError("no products found on the ladders page")
    urls = [SITE_URL + path for path in paths]
    with ThreadPoolExecutor(max_workers=8) as pool:
        pages = list(pool.map(lambda u: _get_json(u, timeout), urls))
    products = tuple(p for p in (_product_from_json(d, u) for d, u in zip(pages, urls)) if p)
    if not products:
        raise ValueError("no priced products found")
    return products


@dataclass
class QuoteLine:
    description: str
    quantity: int
    unit_price: Decimal

    @property
    def amount(self) -> Decimal:
        return (self.unit_price * self.quantity).quantize(CENT, rounding=ROUND_HALF_UP)


@dataclass
class Quote:
    issued_to: str
    issue_date: str
    sales_rep: str
    ship_to: str
    lines: list[QuoteLine] = field(default_factory=list)
    shipping_service: str = DOMESTIC_SERVICE
    shipping_price: Decimal | None = None
    currency: str = "USD"

    @property
    def subtotal(self) -> Decimal:
        return sum((line.amount for line in self.lines), Decimal("0.00"))

    @property
    def excludes_shipping(self) -> bool:
        return bool(self.shipping_service) and self.shipping_price is None

    @property
    def total(self) -> Decimal:
        shipping = self.shipping_price if self.shipping_service and self.shipping_price is not None else Decimal("0.00")
        return self.subtotal + shipping


def validate_quote(quote: Quote) -> list[str]:
    problems = []
    if not quote.issued_to.strip():
        problems.append("Enter who the quote is issued to.")
    if not quote.ship_to.strip():
        problems.append("Enter the shipping address.")
    if not quote.sales_rep.strip():
        problems.append("Enter the sales rep.")
    if not quote.lines:
        problems.append("Add at least one product.")
    for number, line in enumerate(quote.lines, start=1):
        if not line.description.strip():
            problems.append(f"Product {number}: describe the item.")
        if line.quantity < 1:
            problems.append(f"Product {number}: quantity must be at least 1.")
    return problems


def default_file_label(issued_to: str, ship_to: str) -> str:
    skip_line = re.compile(r"^(attn|attention|c/o|dr|mr|mrs|ms|prof)\b\.?", re.IGNORECASE)
    first_line = next((l.strip() for l in ship_to.splitlines() if l.strip() and not skip_line.match(l.strip())), "")
    for source in ((first_line,) if first_line and not first_line[0].isdigit() else ()) + (issued_to,):
        words = [w for w in re.findall(r"[A-Za-z0-9]+", _ascii(source)) if w.lower() not in ("the", "a", "an")]
        if words:
            return words[0].upper()
    return "QUOTE"


def _ascii(text: str) -> str:
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")


def quote_filename(issue_date: str, label: str) -> str:
    safe_label = re.sub(r"[^A-Za-z0-9]+", "_", _ascii(label)).strip("_").upper() or "QUOTE"
    return f"{issue_date.replace('-', '_')}_{safe_label}.docx"


_UNIT_PRICE_WIDTH, _UNIT_PRICE_MIN_PAD = 14, 1
_AMOUNT_WIDTH, _AMOUNT_MIN_PAD = 18, 5
_SHIPPING_EXCLUDED = "{service}\t\t\t\t\t"
_SHIPPING_PRICE_SPACES = 8
_AVENIR_TWIPS = (
    67, 67, 125, 133, 133, 200, 173, 67, 67, 67, 107, 160, 67, 80, 67, 89, 133, 133, 133, 133, 133, 133, 133, 133,
    133, 133, 67, 67, 160, 160, 160, 116, 192, 169, 151, 169, 178, 142, 133, 187, 173, 62, 116, 151, 120, 218, 187, 200,
    142, 200, 147, 133, 138, 169, 147, 231, 151, 143, 138, 62, 89, 62, 160, 120, 58, 125, 147, 116, 147, 133, 76, 147,
    133, 58, 58, 120, 58, 204, 133, 143, 147, 147, 84, 102, 84, 133, 116, 178, 120, 116, 107, 80, 53, 80, 160,
)


def _space_units(text: str) -> int:
    return sum(1 if ch in " .," else 2 for ch in text)


def _text_twips(text: str) -> int:
    return sum(_AVENIR_TWIPS[ord(ch) - 32] if 32 <= ord(ch) < 127 else 133 for ch in text)


def _pad_column(values: list[str], width: int, min_pad: int) -> list[str]:
    target = max([width] + [_space_units(v) + min_pad for v in values])
    return [" " * (target - _space_units(v)) + v for v in values]


def _set_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.text.rstrip(" ") == text.rstrip(" "):
        return
    runs = paragraph.runs
    keep = next((r for r in runs if r.text), runs[0] if runs else None)
    if keep is None:
        paragraph.add_run(text)
        return
    for run in runs:
        if run is not keep:
            run._r.getparent().remove(run._r)
    keep.text = text


def _set_alignment(paragraph: Paragraph, alignment: WD_ALIGN_PARAGRAPH) -> None:
    if paragraph.alignment != alignment:
        paragraph.alignment = alignment


def _insert_after(anchor: Paragraph, prototype: Paragraph, text: str) -> Paragraph:
    new_p = copy.deepcopy(prototype._p)
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    _set_text(paragraph, text)
    return paragraph


def _remove(paragraph: Paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def _clean_lines(text: str) -> list[str]:
    return [line.strip() for line in str(text or "").splitlines() if line.strip()]


def _sync_lines(heading: Paragraph, paragraphs: list[Paragraph], lines: list[str]) -> None:
    for paragraph, line in zip(paragraphs, lines):
        _set_text(paragraph, line)
    if len(lines) > len(paragraphs):
        if paragraphs:
            anchor = prototype = paragraphs[-1]
        else:
            anchor = heading
            prototype = Paragraph(copy.deepcopy(heading._p), heading._parent)
            for run in prototype.runs:
                run.bold = False
        for line in lines[len(paragraphs):]:
            anchor = _insert_after(anchor, prototype, line)
    for paragraph in paragraphs[len(lines):]:
        _remove(paragraph)


def _fill_sender_block(doc: Document, sender: dict[str, str]) -> None:
    paragraphs = doc.paragraphs
    email_at = next((i for i, p in enumerate(paragraphs) if p.text.strip().startswith("e:")), None)
    if email_at is None:
        return
    name_at = next((i for i, p in enumerate(paragraphs[:email_at]) if p.text.strip()), None)
    if name_at is None:
        return
    phone_at = email_at + 1 if email_at + 1 < len(paragraphs) and paragraphs[email_at + 1].text.strip().startswith("t:") else email_at
    _set_text(paragraphs[name_at], sender.get("company_name", ""))
    lines = _clean_lines(sender.get("company_address", "")) + [
        f"e: {sender.get('company_email', '')}",
        f"t: {sender.get('company_phone', '')}",
    ]
    _sync_lines(paragraphs[name_at], paragraphs[name_at + 1 : phone_at + 1], lines)


def _fill_quote_box(doc: Document, quote: Quote) -> bool:
    values = {
        "Quote issued to:": quote.issued_to.strip(),
        "Issue date:": quote.issue_date.strip(),
        "Sales rep:": quote.sales_rep.strip(),
    }
    found = False
    for box in doc.element.body.iter(qn("w:txbxContent")):
        for p in box.iter(qn("w:p")):
            paragraph = Paragraph(p, None)
            text = paragraph.text
            for label, value in values.items():
                if label not in text:
                    continue
                found = True
                if text.strip() == f"{label} {value}":
                    continue
                runs = paragraph.runs
                label_at = next((i for i, r in enumerate(runs) if label in r.text), None)
                if label_at is None:
                    _set_text(paragraph, f"{label} {value}")
                    continue
                value_proto = runs[label_at + 1] if label_at + 1 < len(runs) else None
                new_run = copy.deepcopy((value_proto or runs[label_at])._r)
                for run in runs[label_at + 1 :]:
                    run._r.getparent().remove(run._r)
                runs[label_at]._r.addnext(new_run)
                value_run = Paragraph(p, None).runs[label_at + 1]
                value_run.text = f" {value}"
                if value_proto is None:
                    value_run.bold = False
    return found


def _find_table(doc: Document, test: Callable[[Table], bool]) -> Table | None:
    for table in doc.tables:
        try:
            if test(table):
                return table
        except IndexError:
            continue
    return None


_SHIP_TO_MIN_LINES = 4


def _fill_ship_to(table: Table, ship_to: str) -> None:
    cell = table.cell(0, 0)
    paragraphs = cell.paragraphs
    heading_at = next(i for i, p in enumerate(paragraphs) if "Ship to" in p.text)
    following = paragraphs[heading_at + 1 :]
    last_text = max((i for i, p in enumerate(following) if p.text.strip()), default=-1)
    _sync_lines(paragraphs[heading_at], following[: last_text + 1], _clean_lines(ship_to))

    following = cell.paragraphs[heading_at + 1 :]
    if following:
        anchor = following[-1]
        for _ in range(_SHIP_TO_MIN_LINES - len(following)):
            blank = copy.deepcopy(anchor._p)
            for indent in blank.findall(qn("w:pPr") + "/" + qn("w:ind")):
                indent.getparent().remove(indent)
            anchor._p.addnext(blank)
            anchor = Paragraph(blank, anchor._parent)
            _set_text(anchor, "")


def _set_cell(cell, text: str) -> None:
    paragraphs = cell.paragraphs
    for extra in paragraphs[1:]:
        _remove(extra)
    _set_text(paragraphs[0], text)


def _fill_items(table: Table, quote: Quote) -> None:
    tbl = table._tbl
    body = list(table.rows)[1:]
    if not body:
        prototype = copy.deepcopy(table.rows[0]._tr)
        for tag in ("w:shd", "w:b"):
            for element in list(prototype.iter(qn(tag))):
                element.getparent().remove(element)
        tbl.append(prototype)
        body = list(table.rows)[1:]
    while len(body) < len(quote.lines):
        added = copy.deepcopy(body[-1]._tr)
        for top in list(added.iter(qn("w:top"))):
            borders = top.getparent()
            if borders.tag == qn("w:tcBorders"):
                borders.remove(top)
                if len(borders) == 0:
                    borders.getparent().remove(borders)
        body[-1]._tr.addnext(added)
        body = list(table.rows)[1:]
    for row in body[len(quote.lines):]:
        tbl.remove(row._tr)

    unit_prices = _pad_column([format_money(l.unit_price) for l in quote.lines], _UNIT_PRICE_WIDTH, _UNIT_PRICE_MIN_PAD)
    amounts = _pad_column([format_money(l.amount) for l in quote.lines], _AMOUNT_WIDTH, _AMOUNT_MIN_PAD)
    for row, line, unit_price, amount in zip(list(table.rows)[1:], quote.lines, unit_prices, amounts):
        cells = row.cells
        _set_cell(cells[0], str(line.quantity))
        _set_cell(cells[1], " ".join(line.description.split()))
        _set_cell(cells[2], unit_price)
        _set_cell(cells[3], amount)


def _fill_shipping_line(doc: Document, items: Table, total: Table, quote: Quote) -> None:
    shipping = None
    element = items._tbl.getnext()
    while element is not None and element is not total._tbl:
        if element.tag == qn("w:p") and Paragraph(element, items._parent).text.strip():
            shipping = Paragraph(element, items._parent)
            break
        element = element.getnext()

    if not quote.shipping_service:
        if shipping is not None:
            _remove(shipping)
        return
    if shipping is None:
        new_p = copy.deepcopy(items.rows[-1].cells[1].paragraphs[0]._p)
        items._tbl.addnext(new_p)
        shipping = Paragraph(new_p, items._parent)
        shipping.paragraph_format.left_indent = Twips(1440)
        shipping.paragraph_format.first_line_indent = Twips(720)
        shipping.paragraph_format.space_after = Twips(0)
        shipping.paragraph_format.line_spacing = 1.0

    if quote.shipping_price is None:
        _set_text(shipping, _SHIPPING_EXCLUDED.format(service=quote.shipping_service))
        _set_alignment(shipping, WD_ALIGN_PARAGRAPH.CENTER)
        return

    fmt = shipping.paragraph_format
    start_twips = sum(length.twips for length in (fmt.left_indent, fmt.first_line_indent) if length is not None)
    tab_stop = _default_tab_stop(doc)
    widths = [int(float(g.get(qn("w:w")))) for g in items._tbl.tblGrid.findall(qn("w:gridCol"))]
    last_stop = sum(widths[:3]) // tab_stop * tab_stop
    text_end = start_twips + _text_twips(quote.shipping_service)
    tabs = max(1, last_stop // tab_stop - text_end // tab_stop)

    price = format_money(quote.shipping_price)
    amount_width = max([_AMOUNT_WIDTH] + [_space_units(format_money(l.amount)) + _AMOUNT_MIN_PAD for l in quote.lines])
    spaces = max(1, _SHIPPING_PRICE_SPACES + (amount_width - _AMOUNT_WIDTH) + (_space_units("$000.00") - _space_units(price)))
    _set_text(shipping, quote.shipping_service + "\t" * tabs + " " * spaces + price)
    _set_alignment(shipping, WD_ALIGN_PARAGRAPH.LEFT)


def _default_tab_stop(doc: Document) -> int:
    try:
        value = doc.settings.element.find(qn("w:defaultTabStop")).get(qn("w:val"))
        return int(float(value)) or 720
    except (AttributeError, TypeError, ValueError):
        return 720


def _widen_total_value_column(table: Table, text: str) -> None:
    grid = table._tbl.tblGrid.findall(qn("w:gridCol"))
    if len(grid) != 2:
        return
    label_w, value_w = (int(float(g.get(qn("w:w")))) for g in grid)
    needed = _space_units(text.replace("USD", "")) * 67 + 480 + 216 + 150
    if needed <= value_w:
        return
    shift = min(needed - value_w, label_w // 2)
    grid[0].set(qn("w:w"), str(label_w - shift))
    grid[1].set(qn("w:w"), str(value_w + shift))
    for row in table.rows:
        cells = row.cells
        if len(cells) == 2:
            for cell, width in zip(cells, (label_w - shift, value_w + shift)):
                if cell._tc.tcPr is not None and cell._tc.tcPr.find(qn("w:tcW")) is not None:
                    cell.width = Twips(width)


def _fill_total(table: Table, quote: Quote) -> None:
    label_cell, value_cell = table.rows[0].cells[0], table.rows[0].cells[-1]
    _set_cell(label_cell, " Total:")
    value = f" {format_money(quote.total)} {quote.currency} "
    paragraphs = value_cell.paragraphs
    _set_text(paragraphs[0], value)
    if quote.excludes_shipping:
        if len(paragraphs) >= 2:
            note = paragraphs[1]
            _set_text(note, "(excluding shipping)")
        else:
            note = _insert_after(paragraphs[0], paragraphs[0], "(excluding shipping)")
        _set_alignment(note, WD_ALIGN_PARAGRAPH.CENTER)
        surplus = paragraphs[2:]
    else:
        surplus = paragraphs[1:]
    for paragraph in surplus:
        _remove(paragraph)
    _widen_total_value_column(table, value.strip())


def _fill_payment_lines(doc: Document, sender: dict[str, str]) -> None:
    lines = (
        ("Bank Name:", f"Bank Name: {sender.get('bank_name', '')}"),
        ("Account Number:", f"Account Number: {sender.get('account_number', '')}"),
        ("Routing Number (For Direct", f"Routing Number (For Direct Deposit): {sender.get('direct_deposit_routing', '')}"),
        ("Routing Number (For Wire", f"Routing Number (For Wire Transfer): {sender.get('wire_routing', '')}"),
        ("Swift Code:", f"Swift Code: {sender.get('swift_code', '')}"),
    )
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        for prefix, replacement in lines:
            if text.startswith(prefix):
                _set_text(paragraph, replacement)
                break


def generate_quote_docx(template_bytes: bytes, quote: Quote, sender: dict[str, str]) -> bytes:
    doc = Document(io.BytesIO(template_bytes))
    items = _find_table(doc, lambda t: "Quantity" in t.rows[0].cells[0].text and "Item" in t.rows[0].cells[1].text)
    total = _find_table(doc, lambda t: len(t.rows[0].cells) >= 2 and t.rows[0].cells[0].text.strip().startswith("Total"))
    ship_to = _find_table(doc, lambda t: "Ship to" in t.cell(0, 0).text)
    missing = [name for name, part in (("the item table", items), ("the total", total), ("the Ship to box", ship_to)) if part is None]
    if items is None or total is None or ship_to is None:
        raise QuoteTemplateError(f"This file is not a Simplex quote template. Missing: {', '.join(missing)}.")
    if len(items.columns) < 4:
        raise QuoteTemplateError("The item table needs four columns: Quantity, Item, Unit price, Amount.")
    if not _fill_quote_box(doc, quote):
        raise QuoteTemplateError('This file is not a Simplex quote template. Missing: the "Quote issued to:" box.')

    _fill_sender_block(doc, sender)
    _fill_ship_to(ship_to, quote.ship_to)
    _fill_items(items, quote)
    _fill_shipping_line(doc, items, total, quote)
    _fill_total(total, quote)
    _fill_payment_lines(doc, sender)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
