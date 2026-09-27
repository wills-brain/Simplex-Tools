from __future__ import annotations

import io
import re
import tempfile
import zipfile
from dataclasses import dataclass, asdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pymupdf as fitz
import pandas as pd
import streamlit as st
from docxtpl import DocxTemplate
from docx import Document
from docx.shared import Pt, Inches

import generic_po
import thermo_qty
from quotes import (
    BUNDLED_CATALOG,
    BUNDLED_CATALOG_DATE,
    DOMESTIC_SERVICE,
    FEDEX_RATES_URL,
    INTERNATIONAL_SERVICE,
    SHIPPING_SERVICES,
    Product,
    Quote,
    QuoteLine,
    QuoteTemplateError,
    default_file_label,
    fetch_catalog,
    format_money,
    generate_quote_docx,
    parse_money,
    quote_filename,
    validate_quote,
)

try:
    import pytesseract
    from PIL import Image, ImageOps, ImageFilter, ImageEnhance
except Exception:
    pytesseract = None
    Image = None
    ImageOps = None
    ImageFilter = None
    ImageEnhance = None

APP_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = APP_DIR / "simplex_invoice_template.docx"
DEFAULT_QUOTE_TEMPLATE = APP_DIR / "simplex_quote_template.docx"
BRAND_LOGO = APP_DIR / "simplex_logo_navy.png"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
APP_VERSION = "1.3.9"
APP_PASSCODE = "simplexlovesdads"

MONEY_RE = r"(?:\$\s*)?[0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]{2})|(?:\$\s*)?[0-9]+(?:\.[0-9]{2})"

SENDER_DEFAULTS = {
    "company_name": "Simplex Sciences",
    "company_address": "206 Elm Street\nPO Box 206602\nNew Haven, CT\n06520 United States",
    "company_email": "contact@simplexsciences.com",
    "company_phone": "+1 (510) 847-9188",
    "bank_name": "Wells Fargo",
    "account_number": "7462713764",
    "direct_deposit_routing": "021101108",
    "wire_routing": "121000248",
    "swift_code": "WFBIUS6S",
}

THERMO_BILL_TO_DEFAULT = "Fisher Scientific\nP.O. Box 1768\nPittsburgh, PA 15230\nEmail: APFax@thermofisher.com"


@dataclass
class ParsedPO:
    invoice_issued_to: str = "Fisher Scientific"
    issue_date: str = ""
    order_number: str = ""
    shipping_date: str = ""
    sales_rep: str = ""
    customer_po_number: str = ""
    tracking_number: str = ""
    bill_to: str = ""
    ship_to: str = ""
    total: str = ""
    currency: str = "USD"
    source_file: str = ""
    line_items: list[dict[str, str]] | None = None
    raw_text: str = ""
    ocr_rotation_used: str = ""
    document_kind: str = "thermo"
    notes: list[str] | None = None
    shipping_cost: str = ""


def clean_money(value: Any, with_symbol: bool = True) -> str:
    raw = str(value or "").strip().replace(" ", "")
    if not raw:
        return ""
    raw = raw.replace("$", "")
    try:
        num = float(raw.replace(",", ""))
        return f"${num:,.2f}" if with_symbol else f"{num:.2f}"
    except ValueError:
        return ("$" + raw) if with_symbol and not raw.startswith("$") else raw


def money_to_float(value: Any) -> float:
    raw = str(value or "").replace("$", "").replace(",", "").strip()
    try:
        return float(raw)
    except ValueError:
        return 0.0


def safe_filename(name: str) -> str:
    base = Path(str(name)).stem or "invoice"
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("_")
    return base[:80] or "invoice"


def normalize_ocr_text(text: str) -> str:
    replacements = {
        "Selontitic": "Scientific",
        "SelenBlte": "Scientific",
        "Selentific": "Scientific",
        "Therno": "Thermo",
        "Pittsburgh, PA 18230": "Pittsburgh, PA 15230",
        "Pittsburgh PA 18230": "Pittsburgh, PA 15230",
        "HEW HAVEN": "NEW HAVEN",
        "HET CONTACT": "MET CONTACT",
        "INSTRICTIONS": "INSTRUCTIONS",
        "INSTRICTIONS xxx": "INSTRUCTIONS ***",
        "EMANLORMAN": "EMAIL OR MAIL",
        "APFax@thermofishercom": "APFax@thermofisher.com",
        "APVendor@thermofishercom": "APVendor@thermofisher.com",
        "FO. Box": "P.O. Box",
        "P.O Box": "P.O. Box",
        "PO. Box": "P.O. Box",
        "UREFF": "REF#",
        "UREF#": "REF#",
        "IREF#": "REF#",
        "THC": "INC",
        "GRAIL THC": "GRAIL INC",
        "ss50": "SS50",
        "ss5o": "SS50",
        "DHA LADDER": "DNA LADDER",
        "DHA LADOER": "DNA LADDER",
        "TADDER": "LADDER",
    }
    out = text
    for old, new in replacements.items():
        out = out.replace(old, new)
    out = out.replace("\u2014", "-").replace("\u2013", "-")
    return out


def score_thermo_text(text: str) -> int:
    low = text.lower()
    terms = ["purchase order", "fisher", "scientific", "supplier", "ship to", "invoice", "payable", "quantity", "unit cost", "extended", "dna ladder"]
    return sum(3 if term in low else 0 for term in terms) + min(len(text) // 200, 10)


def extract_text_from_pdf(file_bytes: bytes, enable_ocr: bool = True) -> tuple[str, str]:
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        text_parts = [(page.get_text("text") or "") for page in pdf]
        extracted = "\n".join(text_parts).strip()
        if score_thermo_text(extracted) >= 9:
            return normalize_ocr_text(extracted), "text-layer"

        if not enable_ocr or pytesseract is None or Image is None:
            return normalize_ocr_text(extracted), "text-layer/ocr-off"

        all_page_text = []
        rotations = []
        for page in pdf:
            pix = page.get_pixmap(dpi=250)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            candidates: list[tuple[int, str, int]] = []
            for angle in [0, 90, 180, 270]:
                rotated = img.rotate(angle, expand=True)
                try:
                    txt = pytesseract.image_to_string(rotated, config="--psm 6")
                except Exception:
                    txt = ""
                txt = normalize_ocr_text(txt)
                candidates.append((score_thermo_text(txt), txt, angle))
            score, txt, angle = max(candidates, key=lambda x: x[0])
            all_page_text.append(txt)
            rotations.append(str(angle))
        return "\n".join(all_page_text).strip(), ",".join(rotations)


def extract_filename_po(filename: str) -> str:
    m = re.search(r"(\d{5,})", Path(filename).stem)
    return m.group(1) if m else ""


def find_purchase_order_number(text: str, filename: str) -> str:
    from_filename = extract_filename_po(filename)
    patterns = [
        r"Purchase\s+Order\s+Number\s*[:#-]?\s*([A-Z0-9-]{5,})",
        r"Customer\s+Purchase\s+Order\s+Number\s*[:#-]?\s*([A-Z0-9-]{5,})",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if m:
            candidate = m.group(1).strip()
            if not re.search(r"scientific|order|number", candidate, re.I):
                return candidate
    return from_filename


def clean_order_token(token: str) -> str:
    token = re.sub(r"[^A-Z0-9-]", "", str(token or "").upper())
    return token.strip("-_")


def normalize_fisher_order_candidate(raw: str) -> str:
    raw_text = str(raw or "").upper()
    token = re.sub(r"[^A-Z0-9]", "", raw_text)
    if not token:
        return ""

    if token.startswith(("DBS", "D8S", "D5S")):
        token = "DR5" + token[3:]
    else:
        for prefix in ("DBR", "D8R", "DER", "DPR", "DRS", "D5R"):
            if token.startswith(prefix):
                token = "DR" + token[len(prefix):]
                break
    if not token.startswith("DR"):
        if token.startswith(("DB", "D8", "DS", "D5")):
            token = "DR" + token[2:]
        elif token.startswith("D") and len(token) >= 5 and token[1] not in "R":
            token = "DR" + token[1:]

    if token.startswith("DR"):
        suffix = token[2:].translate(str.maketrans({"O": "0", "Q": "0", "I": "1", "L": "1", "B": "8", "S": "5", "Z": "2", "T": "7"}))
        m = re.search(r"(\d{5,12})", suffix)
        if m:
            return "DR" + m.group(1)
    return ""

def normalize_product_name(description: str) -> str:
    raw = str(description or "").strip()
    if not raw:
        return ""
    upper = raw.upper().replace("$", "S")
    upper = upper.replace("SS5O", "SS50").replace("S5O", "SS50")
    upper = upper.replace("S660", "SS50").replace("S650", "SS50").replace("S560", "SS50")
    upper = upper.replace("SS 50", "SS50").replace("SS 20", "SS20")
    upper = upper.replace("DHA", "DNA").replace("TADDER", "LADDER").replace("LADOER", "LADDER")
    upper = upper.replace("1000L", "100UL").replace("1001L", "100UL").replace("1000IL", "100UL")
    upper = upper.replace("100 U L", "100UL").replace("100 L", "100UL").replace("100ΜL", "100UL")
    m = re.search(r"S{1,2}\s*(\d{2})\s*DNA\s+LADDER(?:\s*(\d+)\s*(?:U|µ)?L)?", upper, re.I)
    if m:
        size = m.group(1)
        volume = m.group(2) or "100"
        return f"ss{size} DNA Ladder ({volume}uL)"
    m2 = re.search(r"(?:660|650|560)\s*DNA\s+LADDER(?:\s*(\d+)\s*(?:U|µ)?L)?", upper, re.I)
    if m2:
        volume = m2.group(1) or "100"
        return f"ss50 DNA Ladder ({volume}uL)"
    cleaned = re.sub(r"\s+", " ", raw).strip()
    cleaned = re.sub(r"\bDNA\s+LADDER\b", "DNA Ladder", cleaned, flags=re.I)
    return cleaned

def find_fisher_scientific_order_number(text: str, filename: str) -> str:
    normalized_text = normalize_ocr_text(text)

    label_match = re.search(r"Fisher\s+Scientific\s+Order\s+Number", normalized_text, re.I)
    if label_match:
        local = normalized_text[label_match.end(): label_match.end() + 550]
        local = re.split(r"Order\s+Date|Quote\s+Number|Terms\s+and\s+Conditions", local, flags=re.I)[0]
        for m in re.finditer(r"\b(?:D\s*[B8EP]?\s*R\s*|D\s*[B8S5]{1,2}\s+)[A-Z0-9]{2,12}(?:\s+[A-Z0-9]{2,8})?\b", local, re.I):
            candidate = normalize_fisher_order_candidate(m.group(0))
            if candidate and candidate.startswith("DR"):
                return candidate
        for m in re.finditer(r"\bD[A-Z0-9]{5,12}(?:\s+[A-Z0-9]{2,8})?\b", local, re.I):
            candidate = normalize_fisher_order_candidate(m.group(0))
            if candidate and candidate.startswith("DR"):
                return candidate

    lines = [clean_ship_line(x) for x in normalized_text.splitlines() if clean_ship_line(x)]
    for i, line in enumerate(lines):
        if re.search(r"Fisher\s+Scientific\s+Order\s+Number", line, re.I):
            window = " ".join(lines[i:i + 5])
            for m in re.finditer(r"\b(?:D\s*[B8EP]?\s*R\s*|D\s*[B8S5]{1,2}\s+)[A-Z0-9]{2,10}(?:\s+[A-Z0-9]{2,8})?\b", window, re.I):
                candidate = normalize_fisher_order_candidate(m.group(0))
                if candidate:
                    return candidate

            for nxt in lines[i + 1:i + 4]:
                tokens = re.findall(r"\b[A-Z0-9][A-Z0-9-]{5,}\b", nxt.upper())
                for token in tokens:
                    token = clean_order_token(token)
                    if token and not re.search(r"ORDER|DATE|QUOTE|NUMBER|CUSTOMER|ACCOUNT|TERMS|FISHER|SCIENTIFIC|THERMO|BRAND", token):
                        candidate = normalize_fisher_order_candidate(token)
                        if candidate:
                            return candidate

    candidates: list[str] = []
    for m in re.finditer(r"\b(?:D\s*[B8EP]?\s*R\s*|D\s*[B8S5]{1,2}\s+)[A-Z0-9]{2,10}(?:\s+[A-Z0-9]{2,8})?\b", normalized_text, re.I):
        cand = normalize_fisher_order_candidate(m.group(0))
        if cand and cand.startswith("DR"):
            candidates.append(cand)
    if candidates:
        return candidates[0]

    return ""


def extract_bill_to(text: str) -> str:
    email = "APFax@thermofisher.com" if re.search(r"ap\s*fax|apfax|apifax", text, re.I) else "APFax@thermofisher.com"
    m = re.search(r"P\.?O\.?\s*Box\s*1768.*?Pittsburgh,?\s*PA\s*15[28]30", text, re.I | re.S)
    if m:
        return f"Fisher Scientific\nP.O. Box 1768\nPittsburgh, PA 15230\nEmail: {email}"
    return THERMO_BILL_TO_DEFAULT


def clean_ship_line(line: str) -> str:
    line = re.sub(r"^[|!\[\](){}/\\\-\s]+", "", line).strip()
    line = re.sub(r"[|!\[\]{}]+$", "", line).strip()
    line = re.sub(r"\s{2,}", " ", line)
    return line


def _looks_like_ship_company(line: str) -> bool:
    return bool(re.search(r"\b(GRAIL|MAYO|CLINIC|SCIEX|UNIVERSITY|HOSPITAL|LAB|LABORATORY|INC|LLC|CORP|COMPANY)\b", line, re.I))


def _format_ship_line(line: str) -> str:
    ln = clean_ship_line(line)
    ln = ln.replace("REF¥", "REF#").replace("REFY", "REF#").replace("REFF", "REF#")
    ln = ln.replace("GRAIL THC", "GRAIL INC").replace("GRAIL THG", "GRAIL INC")
    ln = ln.replace("HC $4", "NC 54").replace("HC 54", "NC 54").replace("NC $4", "NC 54")
    ln = re.split(r"\s+i\s+PURCHASING\b|\s+PURCHASING\b|\s+i\s+SCOTT\b|\s+SCOTT\s+MORES\b|\s+Pittsburgh\b", ln, flags=re.I)[0]
    ln = re.sub(r"\bS5901\b", "55901", ln)
    ln = re.sub(r"\bVEST\b", "WEST", ln, flags=re.I)
    ln = re.sub(r"\bASSEMBLY\b.*$", "ASSEMBLY", ln, flags=re.I)
    ln = re.sub(r"\s+L$", "", ln)
    ln = re.sub(r"\bROCHESTER\s+MN\.?\s*(\d{5})", r"Rochester, MN \1", ln, flags=re.I)
    ln = re.sub(r"\bDURHAM\s+NC\.?\s*(\d{5})", r"Durham, NC \1", ln, flags=re.I)
    ln = re.sub(r"\bMARLBOROUGH,?\s+MASSACHUSETTS\b", "Marlborough, Massachusetts", ln, flags=re.I)
    ln = re.sub(r"^REF\s*[#:]?\s*[:#-]*\s*", "Ref: ", ln, flags=re.I)
    ln = re.sub(r"^Ref:\s*[:#-]+\s*", "Ref: ", ln, flags=re.I)
    ln = re.sub(r"^ATTN\s*[:\-]?\s*", "ATTN: ", ln, flags=re.I)
    ln = re.sub(r"\s+", " ", ln).strip(" |-()")
    return ln


def extract_ship_to(text: str) -> str:
    raw_lines = [x for x in text.splitlines() if x.strip()]
    cleaned: list[str] = []

    for raw in raw_lines:
        ln = clean_ship_line(raw)
        low = ln.lower()

        if re.search(r"order notes|product description|catalog|total:|page:", ln, re.I):
            break
        if re.search(r"purchase order|fisher scientific order|order date|terms and conditions|supplier number|customer purchase|msds|caller:|due date", ln, re.I):
            continue

        if "|" in ln:
            parts = [clean_ship_line(p) for p in ln.split("|")]
            candidates = []
            for part in parts:
                if not part:
                    continue
                if re.search(r"SIMPLEX", part, re.I) and _looks_like_ship_company(part):
                    part = re.sub(r"^.*?SIMPLEX\s+SCIENCES\s*['iI,;:-]*\s*", "", part, flags=re.I)
                if re.search(r"EDWARDS|NEW HAVEN|Fisher Scientific\s*$|Pittsburgh|APVendor|APFax|Telephone|SCOTT MORES|PURCHASING|Address:\s*Fisher|Email:", part, re.I):
                    trimmed = re.split(r"\s+Pittsburgh\b|\s+i\s+PURCHASING|\s+PURCHASING", part, flags=re.I)[0]
                    if trimmed != part and re.search(r"DURHAM|ROCHESTER|MARLBOROUGH|\b\d{3,5}\b.*\b(HWY|HIGHWAY|ROAD|RD|STREET|ST|AVE|WAY|DRIVE|DR|WEST|EAST|NORTH|SOUTH|ASSEMBLY)\b", trimmed, re.I):
                        candidates.append(trimmed)
                    continue
                candidates.append(part)
            for part in candidates:
                if re.search(r"REF[#¥YF]?|ATTN|\b\d{3,5}\b.*\b(HWY|ROAD|RD|ST|STREET|AVE|WAY|DR|DRIVE|WEST|EAST|NORTH|SOUTH|ASSEMBLY)\b|\b(GRAIL|MAYO|CLINIC|SCIEX|INC|LLC)\b|\b(DURHAM|ROCHESTER|MARLBOROUGH)\b", part, re.I):
                    cleaned.append(_format_ship_line(part))
                    break
            continue

        if _looks_like_ship_company(ln) and not re.search(r"SIMPLEX|Fisher Scientific Company", ln, re.I):
            part = re.sub(r".*?(GRAIL\s+INC|MAYO\s+CLINIC[^|]*|AB\s+SCIEX\s+LLC).*", r"\1", ln, flags=re.I)
            cleaned.append(_format_ship_line(part))
            continue

        if re.search(r"REF[#¥YF]", ln, re.I):
            part = re.sub(r".*?(REF[#¥YF]?\s*[:#-]?\s*[A-Z0-9-]+)", r"\1", ln, flags=re.I)
            part = re.split(r"\s+P\.?O\.?\s*Box|\s+Pittsburgh|\|", part, flags=re.I)[0]
            cleaned.append(_format_ship_line(part))
            continue

        if re.search(r"\b\d{3,5}\b.*\b(HWY|HIGHWAY|ROAD|RD|STREET|ST|AVE|WAY|DRIVE|DR|WEST|EAST|NORTH|SOUTH|ASSEMBLY)\b", ln, re.I):
            part = ln
            part = re.sub(r"^.*?((?:\d{3,5})\s+[^|]*(?:HWY|HIGHWAY|ROAD|RD|STREET|ST|AVE|WAY|DRIVE|DR|WEST|EAST|NORTH|SOUTH|ASSEMBLY)[^|]*)", r"\1", part, flags=re.I)
            part = re.split(r"\s+Pittsburgh\b|\s+PURCHASING\b|\|", part, flags=re.I)[0]
            cleaned.append(_format_ship_line(part))
            continue

        if re.search(r"\b(DURHAM|ROCHESTER|MARLBOROUGH)\b", ln, re.I):
            part = re.sub(r".*?\b(DURHAM\s+NC\.?\s*\d{5}|ROCHESTER\s+MN\.?\s*\S{5}|MARLBOROUGH,?\s+MASSACHUSETTS).*", r"\1", ln, flags=re.I)
            cleaned.append(_format_ship_line(part))
            continue

        if re.search(r"\bATTN\b", ln, re.I):
            part = re.sub(r".*?(ATTN\s*[:\-]?\s*[^|]+).*", r"\1", ln, flags=re.I)
            cleaned.append(_format_ship_line(part))
            continue

    if len(cleaned) < 3:
        lines = [clean_ship_line(x) for x in text.splitlines()]
        start = None
        for i, ln in enumerate(lines):
            if re.search(r"Ship\s*To\s*Information|Shi\s*To\s*Information", ln, re.I):
                start = i + 1
                break
        if start is not None:
            for ln in lines[start:start + 12]:
                if re.search(r"send\s*invoice|accounts payable|purchasing agent|customer purchase|supplier number|order notes|telephone|email:", ln, re.I):
                    break
                mid = ln
                if "|" in mid:
                    pieces = [clean_ship_line(p) for p in mid.split("|") if clean_ship_line(p)]
                    if len(pieces) >= 2:
                        mid = pieces[1]
                mid = re.sub(r"\bEmail:.*$|\bAddress:.*$", "", mid, flags=re.I).strip()
                if mid and not re.search(r"SIMPLEX|EDWARDS|NEW HAVEN|APVendor|APFax|Fisher Scientific\s*$", mid, re.I):
                    cleaned.append(_format_ship_line(mid))

    final: list[str] = []
    for ln in cleaned:
        ln = _format_ship_line(ln)
        if not ln:
            continue
        if re.search(r"SIMPLEX|EDWARDS|NEW HAVEN|Pittsburgh|APVendor|APFax|SCOTT MORES|PURCHASING|Telephone|Supplier information|Ship To Information", ln, re.I):
            continue
        if ln.lower() not in {x.lower() for x in final}:
            final.append(ln)

    if final and not any(re.search(r"United States", x, re.I) for x in final):
        final.append("United States of America")
    return "\n".join(final)

def parse_line_items(text: str) -> list[dict[str, str]]:
    lines = [clean_ship_line(x) for x in text.splitlines() if clean_ship_line(x)]
    items: list[dict[str, str]] = []

    for line in lines:
        low = line.lower()
        if not ("ladder" in low or "dna" in low or "ss50" in low or "ss20" in low or "s660" in low):
            continue

        qty = ""
        m_qty = re.search(r"\|?\s*(\d+)\s*\|?\s*(?:EA|EACH|PK|CS)\b", line, re.I)
        if m_qty:
            qty = m_qty.group(1)

        money_values = re.findall(r"\d{1,3}(?:,\d{3})*\.\d{2}", line)
        extended_cost = money_values[-1] if money_values else ""
        unit_cost = ""
        if len(money_values) >= 2:
            unit_cost = money_values[-2]
        elif extended_cost and qty:
            unit_cost = f"{money_to_float(extended_cost) / max(float(qty), 1):.2f}"

        desc_match = re.search(r"((?:SS\s*\d+|S660|S650|SS5O|SSS0)\s*DNA\s+LADDER\s+[^|\n]*)", line, re.I)
        desc = desc_match.group(1) if desc_match else line
        desc = re.split(r"\s+\d+\s*(?:EA|EACH|PK|CS)\b", desc, flags=re.I)[0]
        desc = re.sub(r"\d{1,3}(?:,\d{3})*\.\d{2}.*$", "", desc).strip()
        desc = re.sub(r"\s+", " ", desc).strip(" |-")
        desc = desc.upper()
        desc = desc.replace("SS5O", "SS50").replace("SS 50", "SS50").replace("SS 20", "SS20")
        desc = desc.replace("S660 DNA LADDER", "SS50 DNA LADDER").replace("S650 DNA LADDER", "SS50 DNA LADDER")
        desc = desc.replace("1000L", "100UL").replace("1001L", "100UL").replace("100 U L", "100UL")
        desc = desc.replace("100 L", "100UL")

        if desc and (extended_cost or unit_cost or qty):
            items.append({
                "quantity": qty,
                "item": normalize_product_name(desc),
                "unit_price": clean_money(unit_cost) if unit_cost else "",
                "amount": clean_money(extended_cost) if extended_cost else "",
            })

    if items:
        return items

    for line in lines:
        nums = re.findall(r"\d{1,3}(?:,\d{3})*\.\d{2}", line)
        if nums and re.search(r"\b\d+\s*(EA|EACH|PK|CS)\b", line, re.I):
            qty_match = re.search(r"\b(\d+)\s*(EA|EACH|PK|CS)\b", line, re.I)
            qty = qty_match.group(1) if qty_match else ""
            amount = nums[-1]
            unit = nums[-2] if len(nums) >= 2 else (f"{money_to_float(amount) / max(float(qty or 1), 1):.2f}" if qty else "")
            desc = re.sub(r"\d{1,3}(?:,\d{3})*\.\d{2}.*$", "", line).strip(" |-")
            desc = re.sub(r"^.*?\b([A-Z]{1,4}\d{2,}.*)$", r"\1", desc).strip()
            items.append({"quantity": qty, "item": normalize_product_name(desc), "unit_price": clean_money(unit), "amount": clean_money(amount)})
            break

    return items or [{"quantity": "", "item": "", "unit_price": "", "amount": ""}]


@st.cache_resource(show_spinner=False)
def _easyocr_reader():
    try:
        import easyocr
        return easyocr.Reader(["en"], gpu=False, verbose=False)
    except Exception:
        return None


def _find_label_box(img) -> tuple[int, int, int, int] | None:
    if pytesseract is None or Image is None:
        return None
    try:
        data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT, config="--psm 6")
    except Exception:
        return None
    tokens = data.get("text") or []
    n = len(tokens)
    expected = ["f", "s", "o", "n"]
    best: tuple[int, int, int, int] | None = None
    for i in range(n):
        t = (tokens[i] or "").strip()
        if not t or not t.lower().startswith("fisher"):
            continue
        row_y = data["top"][i]
        row_h = max(1, data["height"][i])
        window = [(i, t)]
        j = i + 1
        while j < n and len(window) < 8:
            tj = (tokens[j] or "").strip()
            if tj:
                if abs(data["top"][j] - row_y) <= row_h * 0.8:
                    window.append((j, tj))
                else:
                    break
            j += 1
        if len(window) < 4:
            continue
        first_letters = [w[1][0].lower() for w in window[:4]]
        if first_letters == expected:
            idxs = [w[0] for w in window[:4]]
            xs = [data["left"][k] for k in idxs]
            ys = [data["top"][k] for k in idxs]
            ws = [data["width"][k] for k in idxs]
            hs = [data["height"][k] for k in idxs]
            x0 = min(xs)
            y0 = min(ys)
            x1 = max(x + w for x, w in zip(xs, ws))
            y1 = max(y + h for y, h in zip(ys, hs))
            if best is None or y0 < best[1]:
                best = (x0, y0, x1, y1)
    return best


def extract_dr_with_easyocr(file_bytes: bytes) -> str:
    if Image is None or ImageOps is None:
        return ""
    reader = _easyocr_reader()
    if reader is None:
        return ""

    try:
        import numpy as np
    except Exception:
        return ""

    best_candidate = ""
    best_score = -1.0

    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        if len(pdf) == 0:
            return ""
        page = pdf[0]
        pix = page.get_pixmap(dpi=500)
        base = Image.open(io.BytesIO(pix.tobytes("png")))

    for angle in (0, 90, 180, 270):
        img = base.rotate(angle, expand=True) if angle else base
        box = _find_label_box(img)
        if box is None:
            continue
        x0, y0, x1, y1 = box
        label_h = max(1, y1 - y0)
        label_w = max(1, x1 - x0)
        crop_x0 = max(0, int(x0 - label_w * 0.10))
        crop_x1 = min(img.size[0], int(x0 + label_w * 2.0))
        crop_y0 = min(img.size[1], int(y1 + label_h * 0.2))
        crop_y1 = min(img.size[1], int(y1 + label_h * 5.0))
        if crop_y1 <= crop_y0 or crop_x1 <= crop_x0:
            continue
        crop = img.crop((crop_x0, crop_y0, crop_x1, crop_y1))
        arr = np.array(crop)
        try:
            results = reader.readtext(
                arr,
                detail=1,
                paragraph=False,
                allowlist="ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-",
            )
        except Exception:
            results = []
        for entry in results:
            if not entry or len(entry) < 3:
                continue
            _, text, conf = entry[0], entry[1], float(entry[2])
            if conf < 0.45:
                continue
            candidate = normalize_fisher_order_candidate(str(text or ""))
            if not candidate or not candidate.startswith("DR"):
                continue
            suffix_digits = sum(1 for ch in candidate[2:] if ch.isdigit())
            if suffix_digits < 5:
                continue
            score = conf * 10 + suffix_digits
            if score > best_score:
                best_score = score
                best_candidate = candidate
        if best_candidate:
            break

    return best_candidate


def parse_thermofisher_po(text: str, filename: str, file_bytes: bytes | None = None, rotation: str = "") -> ParsedPO:
    text = normalize_ocr_text(text)
    purchase_order_number = find_purchase_order_number(text, filename)
    fisher_order_number = "DR"
    ship_to = extract_ship_to(text)
    bill_to = extract_bill_to(text)
    items = parse_line_items(text)
    subtotal = sum(money_to_float(row.get("amount")) for row in items)
    return ParsedPO(
        invoice_issued_to="Fisher Scientific",
        order_number=purchase_order_number,
        customer_po_number=fisher_order_number,
        tracking_number="",
        bill_to=bill_to,
        ship_to=ship_to,
        total=clean_money(subtotal) if subtotal else "",
        currency="USD",
        source_file=filename,
        line_items=items,
        raw_text=text,
    )


def is_thermo_fisher_document(text: str) -> bool:
    if len(re.sub(r"\s", "", text)) < 40:
        return True
    return bool(re.search(r"fisher|thermo|apfax", text, re.I))


def parse_other_document(file_bytes: bytes, filename: str, text: str, rotation: str, use_ocr: bool) -> ParsedPO:
    fields = generic_po.parse_document(file_bytes, use_ocr=use_ocr, rotation=rotation if rotation[:1].isdigit() else "")
    items = [{**row, "item": normalize_product_name(row["item"]) if row["item"] else ""} for row in fields["line_items"]]
    return ParsedPO(
        invoice_issued_to=fields["invoice_issued_to"],
        customer_po_number=fields["customer_po_number"],
        bill_to=fields["bill_to"],
        ship_to=fields["ship_to"],
        currency=fields["currency"],
        source_file=filename,
        line_items=items,
        raw_text=fields["raw_text"] or text,
        document_kind="other",
        notes=fields["notes"],
        shipping_cost=fields["shipping_cost"],
    )


def fill_missing_quantity(parsed: ParsedPO, text: str, file_bytes: bytes, rotation: str) -> None:
    items = parsed.line_items or []
    if len(items) != 1:
        return
    item = items[0]
    quantity = str(item.get("quantity") or "").strip() or thermo_qty.vote_quantity(text, file_bytes, rotation)
    if not quantity:
        return
    item["quantity"] = quantity
    unit = thermo_qty.unit_price_for(text, quantity)
    if unit is not None and item.get("amount"):
        item["unit_price"] = clean_money(unit)


@st.cache_data(show_spinner=False, max_entries=64)
def read_purchase_order(file_bytes: bytes, filename: str, use_ocr: bool) -> tuple[str, str, ParsedPO]:
    text, rotation = extract_text_from_pdf(file_bytes, enable_ocr=use_ocr)
    if is_thermo_fisher_document(text):
        parsed = parse_thermofisher_po(text, filename, file_bytes=file_bytes, rotation=rotation)
        fill_missing_quantity(parsed, text, file_bytes, rotation)
        return text, rotation, parsed
    parsed = parse_other_document(file_bytes, filename, text, rotation, use_ocr)
    return parsed.raw_text, rotation, parsed


def docx_contains_jinja(template_bytes: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(template_bytes)) as zf:
            xml = "\n".join(zf.read(name).decode("utf-8", errors="ignore") for name in zf.namelist() if name.startswith("word/") and name.endswith(".xml"))
        return ("{{" in xml) or ("{%" in xml)
    except Exception:
        return False


def _set_rfonts(run, font_name: str = "Avenir") -> None:
    run.font.name = font_name
    run.font.size = Pt(12)
    try:
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.rFonts
        if rFonts is None:
            from docx.oxml import OxmlElement
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rFonts.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}", font_name)
    except Exception:
        pass


def _set_paragraph_text(paragraph, text: str, bold: bool | None = None) -> None:
    while len(paragraph.runs) > 1:
        paragraph._p.remove(paragraph.runs[-1]._r)
    if paragraph.runs:
        paragraph.runs[0].text = str(text or "")
        _set_rfonts(paragraph.runs[0])
        if bold is not None:
            paragraph.runs[0].bold = bold
    else:
        run = paragraph.add_run(str(text or ""))
        _set_rfonts(run)
        if bold is not None:
            run.bold = bold


def _set_cell_text(cell, text: str, bold: bool | None = None) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph()
        paragraphs = cell.paragraphs
    _set_paragraph_text(paragraphs[0], str(text or ""), bold=bold)
    try:
        paragraphs[0].paragraph_format.space_after = Pt(0)
        paragraphs[0].paragraph_format.space_before = Pt(0)
    except Exception:
        pass
    for p in list(cell.paragraphs)[1:]:
        _remove_paragraph(p)


def _set_cell_label_body(cell, label: str, body: str, label_bold: bool = True, body_bold: bool = False) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph()
        paragraphs = cell.paragraphs
    p = paragraphs[0]
    for run in list(p.runs):
        p._p.remove(run._r)
    label_run = p.add_run(str(label or ""))
    _set_rfonts(label_run)
    label_run.bold = label_bold
    body_run = p.add_run("\n" + str(body or ""))
    _set_rfonts(body_run)
    body_run.bold = body_bold
    for extra in paragraphs[1:]:
        _set_paragraph_text(extra, "", bold=body_bold)

def _ensure_cell_paragraphs(cell, count: int):
    while len(cell.paragraphs) < count:
        cell.add_paragraph()
    return cell.paragraphs


def _write_para(paragraph, text: str, *, bold: bool | None = None, size_pt: int | float = 12, keep_indent: bool = True) -> None:
    while len(paragraph.runs) > 1:
        paragraph._p.remove(paragraph.runs[-1]._r)
    if paragraph.runs:
        run = paragraph.runs[0]
        run.text = str(text or "")
    else:
        run = paragraph.add_run(str(text or ""))
    _set_rfonts(run)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def _write_label_value_para(paragraph, label: str, value: str, *, label_bold: bool = False, value_bold: bool = False, size_pt: int | float = 12) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    r1 = paragraph.add_run(label)
    _set_rfonts(r1); r1.font.size = Pt(size_pt); r1.bold = label_bold
    r2 = paragraph.add_run(str(value or ""))
    _set_rfonts(r2); r2.font.size = Pt(size_pt); r2.bold = value_bold


def _remove_paragraph(paragraph) -> None:
    try:
        p = paragraph._element
        p.getparent().remove(p)
        paragraph._p = paragraph._element = None
    except Exception:
        pass


def _tidy_para_spacing(paragraph) -> None:
    try:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
    except Exception:
        pass


def _set_left_indent(paragraph, inches: float) -> None:
    try:
        paragraph.paragraph_format.left_indent = Inches(inches)
        paragraph.paragraph_format.first_line_indent = Inches(0)
    except Exception:
        pass


def _write_multiline_body(cell, label: str, body: str) -> None:
    lines = [x.rstrip() for x in str(body or "").splitlines() if x.rstrip()]
    paragraphs = _ensure_cell_paragraphs(cell, max(1 + len(lines), 1))
    _write_para(paragraphs[0], label, bold=True, size_pt=12)
    _tidy_para_spacing(paragraphs[0])
    for i, line in enumerate(lines, start=1):
        _write_para(paragraphs[i], line, bold=False, size_pt=12)
        _tidy_para_spacing(paragraphs[i])
    for p in list(cell.paragraphs)[1 + len(lines):]:
        _remove_paragraph(p)

def _force_avenir_12(doc: Document) -> None:
    for p in doc.paragraphs:
        for r in p.runs:
            current_size = r.font.size
            _set_rfonts(r)
            if current_size is not None:
                r.font.size = current_size
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        current_size = r.font.size
                        _set_rfonts(r)
                        if current_size is not None:
                            r.font.size = current_size


def generate_from_filled_template(template_bytes: bytes, context: dict[str, Any]) -> bytes:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        f.write(template_bytes)
        temp_path = Path(f.name)
    doc = Document(str(temp_path))

    rows = list(context.get("line_items") or [])
    for row in rows:
        if str(row.get("item", "")).strip().lower() != "fedex 2day shipping":
            row["item"] = normalize_product_name(row.get("item", ""))

    if len(doc.tables) >= 1:
        t = doc.tables[0]
        try:
            left = _ensure_cell_paragraphs(t.cell(0, 0), 7)
            _write_para(left[0], context.get('company_name',''), bold=True, size_pt=12)
            addr_lines = str(context.get('company_address','')).splitlines()
            while len(addr_lines) < 4:
                addr_lines.append("")
            _write_para(left[1], addr_lines[0], bold=False, size_pt=12)
            _write_para(left[2], addr_lines[1], bold=False, size_pt=12)
            _write_para(left[3], addr_lines[2], bold=False, size_pt=12)
            _write_para(left[4], addr_lines[3], bold=False, size_pt=12)
            _write_para(left[5], f"e: {context.get('company_email','')}", bold=False, size_pt=12)
            _write_para(left[6], f"t: {context.get('company_phone','')}", bold=False, size_pt=12)
            for p in left[1:7]:
                _set_left_indent(p, 0.18)
                _tidy_para_spacing(p)
            _tidy_para_spacing(left[0])
            for p in left[7:]:
                _write_para(p, "", bold=False, size_pt=12)

            right = _ensure_cell_paragraphs(t.cell(0, 1), 6)
            _write_para(right[0], "Invoice", bold=False, size_pt=20)
            _write_label_value_para(right[1], "Invoice issued to: ", context.get('invoice_issued_to',''), label_bold=False, value_bold=False, size_pt=12)
            _write_label_value_para(right[2], "Issue date: ", context.get('issue_date',''), label_bold=False, value_bold=False, size_pt=12)
            _write_label_value_para(right[3], "Order number: ", context.get('order_number',''), label_bold=False, value_bold=False, size_pt=12)
            _write_label_value_para(right[4], "Shipping date: ", context.get('shipping_date',''), label_bold=False, value_bold=False, size_pt=12)
            _write_label_value_para(right[5], "Sales rep: ", context.get('sales_rep',''), label_bold=False, value_bold=False, size_pt=12)
            for p in right[6:]:
                _write_para(p, "", bold=False, size_pt=12)

            _write_label_value_para(t.cell(1, 0).paragraphs[0], "Customer PO Number: ", context.get('customer_po_number',''), label_bold=True, value_bold=True, size_pt=12)
            try:
                _write_label_value_para(t.cell(1, 1).paragraphs[0], "FedEx Tracking Number: ", context.get('tracking_number',''), label_bold=True, value_bold=True, size_pt=12)
            except Exception:
                pass
            _write_multiline_body(t.cell(2, 0), "Bill to:", context.get('bill_to',''))
            _write_multiline_body(t.cell(2, 1), "Ship to:", context.get('ship_to',''))
        except Exception:
            pass

    if len(doc.tables) >= 2:
        t = doc.tables[1]
        product_rows = [r for r in rows if str(r.get("item", "")).strip().lower() != "fedex 2day shipping"]
        shipping_rows = [r for r in rows if str(r.get("item", "")).strip().lower() == "fedex 2day shipping"]
        product = product_rows[0] if product_rows else {"quantity":"", "item":"", "unit_price":"", "amount":""}
        shipping = shipping_rows[0] if shipping_rows else None
        item_text = f"{product.get('item','')}"
        amount_text = f"{product.get('amount','')}"
        unit_text = f"{product.get('unit_price','')}"
        if shipping and money_to_float(shipping.get("amount")) > 0:
            item_text += f"\n\nFedEx 2Day Shipping"
            amount_text += f"\n\n{shipping.get('amount','')}"
        try:
            _set_cell_text(t.cell(1, 0), f"{product.get('quantity','')}")
            _set_cell_text(t.cell(1, 1), item_text)
            _set_cell_text(t.cell(1, 2), unit_text)
            _set_cell_text(t.cell(1, 3), amount_text)
            last = len(t.rows) - 1
            _set_cell_text(t.cell(last, 2), "  Total:")
            _set_cell_text(t.cell(last, 3), f"{context.get('total','')} {context.get('currency','USD')}")
        except Exception:
            pass

    payment_lines = [
        "Make payments to:",
        f"Bank Name: {context.get('bank_name','')}",
        f"Account Number: {context.get('account_number','')}",
        f"Routing Number (For Direct Deposit): {context.get('direct_deposit_routing','')}",
        f"Routing Number (For Wire Transfer): {context.get('wire_routing','')}",
        f"Swift Code: {context.get('swift_code','')}",
    ]
    pi = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if pi < len(payment_lines) and (txt.startswith("Make payments") or txt.startswith("Bank Name:") or txt.startswith("Account Number:") or txt.startswith("Routing Number") or txt.startswith("Swift Code:")):
            _set_paragraph_text(p, payment_lines[pi])
            pi += 1

    _force_avenir_12(doc)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def generate_docx(template_bytes: bytes, context: dict[str, Any]) -> bytes:
    if not docx_contains_jinja(template_bytes):
        return generate_from_filled_template(template_bytes, context)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        f.write(template_bytes)
        temp_path = Path(f.name)
    tpl = DocxTemplate(str(temp_path))
    tpl.render(context)
    output = io.BytesIO()
    tpl.save(output)
    doc = Document(io.BytesIO(output.getvalue()))
    _force_avenir_12(doc)
    final = io.BytesIO()
    doc.save(final)
    return final.getvalue()

def build_context(sender: dict[str, str], invoice: dict[str, Any], shipping_cost: str) -> dict[str, Any]:
    rows = list(invoice.get("line_items") or [])
    for row in rows:
        if str(row.get("item", "")).strip().lower() != "fedex 2day shipping":
            row["item"] = normalize_product_name(row.get("item", ""))
    ship_cost_float = money_to_float(shipping_cost)
    if ship_cost_float > 0:
        rows.append({"quantity": "", "item": "FedEx 2Day Shipping", "unit_price": "", "amount": clean_money(ship_cost_float)})
    subtotal = sum(money_to_float(r.get("amount")) for r in rows)
    invoice = {**invoice, "line_items": rows, "total": clean_money(subtotal)}
    return {**sender, **invoice}


BRAND_CSS = APP_DIR / "simplex.css"


def apply_brand_style() -> None:
    st.html(BRAND_CSS)


def render_product_header() -> None:
    st.title("Invoices and Quotes", anchor=False)


def render_footer() -> None:
    st.html(f'<div class="ss-footer">Simplex Sciences | Internal use only | Version {APP_VERSION}</div>')


def require_passcode() -> None:
    if st.session_state.get("simplex_passcode_ok"):
        return
    st.markdown("For Simplex Sciences staff. Enter the operations passcode to continue.")
    with st.form("simplex-passcode-form", clear_on_submit=False, border=False, width=420):
        entered = st.text_input("Passcode", type="password", key="simplex_passcode_input")
        submitted = st.form_submit_button("Continue", type="primary")
    if submitted:
        if entered == APP_PASSCODE:
            st.session_state["simplex_passcode_ok"] = True
            st.rerun()
        else:
            st.error("Incorrect passcode.")
    st.stop()


INVOICE_FIELD_LABELS = {
    "invoice_issued_to": "Invoice issued to",
    "customer_po_number": "Customer PO number",
    "bill_to": "Bill to",
    "ship_to": "Ship to",
    "currency": "Currency",
}

SENDER_LABELS = {
    "company_name": "Company name",
    "company_address": "Address",
    "company_email": "Email",
    "company_phone": "Phone",
    "bank_name": "Bank",
    "account_number": "Account number",
    "direct_deposit_routing": "Direct deposit routing number",
    "wire_routing": "Wire routing number",
    "swift_code": "SWIFT code",
}


def main() -> None:
    st.set_page_config(page_title="Invoices and Quotes | Simplex Sciences", page_icon=str(BRAND_LOGO), layout="wide")
    apply_brand_style()
    render_product_header()
    require_passcode()

    with st.sidebar:
        if BRAND_LOGO.exists():
            st.image(str(BRAND_LOGO), width=112)
        st.subheader("Invoice defaults", anchor=False)
        invoice_date_obj = st.date_input("Invoice date", value=date.today())
        sales_rep_default = st.text_input("Sales rep", value="William Shiu")
        st.caption("Used for every purchase order. Shipping date, order number, shipping cost, and tracking number are set for each PDF.")
        uploaded_template = st.file_uploader("Invoice template (.docx)", type=["docx"])
        st.caption("Optional. Without an upload, the app uses simplex_invoice_template.docx from the app folder.")
        use_ocr = st.checkbox("Read scanned PDFs with OCR", value=True)
        sender = dict(SENDER_DEFAULTS)
        with st.expander("Sender and payment details"):
            for key, default in SENDER_DEFAULTS.items():
                label = SENDER_LABELS.get(key, key.replace("_", " ").capitalize())
                if "address" in key:
                    sender[key] = st.text_area(label, value=default, key=f"sender-{key}")
                else:
                    sender[key] = st.text_input(label, value=default, key=f"sender-{key}")
            st.caption("Printed on invoices and quotes.")

    invoice_tab, quote_tab = st.tabs(["Thermo Fisher invoices", "Quotes"])
    with invoice_tab:
        render_invoice_tab(uploaded_template, use_ocr, invoice_date_obj, sales_rep_default, sender)
    with quote_tab:
        render_quote_tab(sender, sales_rep_default)
    render_footer()


def render_invoice_tab(uploaded_template, use_ocr: bool, invoice_date_obj: date, sales_rep_default: str, sender: dict[str, str]) -> None:
    pdf_files = st.file_uploader("Thermo Fisher purchase orders (PDF)", type=["pdf"], accept_multiple_files=True)

    if not pdf_files:
        st.caption("Upload one or more purchase orders. Each one becomes a Word invoice to review before sending.")
        return

    reviewed_ack = st.checkbox("I will check each Word invoice before sending it. OCR can misread values.")
    if not reviewed_ack:
        st.caption("Tick the box above to enable downloads.")
    template_bytes = uploaded_template.getvalue() if uploaded_template else DEFAULT_TEMPLATE.read_bytes()
    all_generated: dict[str, bytes] = {}

    overall_progress = st.progress(0.0, text="Reading purchase orders")
    total_pdfs = len(pdf_files)

    for idx, pdf in enumerate(pdf_files):
        st.divider()
        st.subheader(pdf.name, anchor=False)
        file_bytes = pdf.getvalue()
        overall_progress.progress(idx / max(total_pdfs, 1), text=f"Reading {idx + 1} of {total_pdfs}: {pdf.name}")
        with st.spinner(f"Reading {pdf.name}"):
            try:
                text, rotation, parsed = read_purchase_order(file_bytes, pdf.name, use_ocr)
            except Exception as exc:
                st.error(f"Could not read {pdf.name}: {exc}")
                continue
        is_thermo = parsed.document_kind == "thermo"
        data = asdict(parsed)
        data["ocr_rotation_used"] = rotation
        data["issue_date"] = invoice_date_obj.isoformat()
        data["sales_rep"] = sales_rep_default

        u1, u2, u3, u4 = st.columns(4)
        with u1:
            per_shipping_date_obj = st.date_input("Shipping date", value=date.today(), key=f"shipping-date-{idx}")
        with u2:
            per_internal_order_number = st.text_input("Internal order number", value="FS", key=f"internal-order-{idx}")
        with u3:
            per_shipping_cost = st.text_input("Shipping cost (USD)", value=parsed.shipping_cost, key=f"shipping-cost-{idx}")
        with u4:
            per_tracking_number = st.text_input("FedEx tracking number", value="", key=f"tracking-number-{idx}")

        data["shipping_date"] = per_shipping_date_obj.isoformat()
        data["order_number"] = per_internal_order_number
        data["tracking_number"] = per_tracking_number

        col1, col2 = st.columns([0.95, 1.05])
        with col1:
            st.markdown("##### Invoice fields")
            edited = {}
            for key in ["invoice_issued_to", "customer_po_number", "bill_to", "ship_to", "currency"]:
                label = INVOICE_FIELD_LABELS[key]
                if key in {"bill_to", "ship_to"}:
                    edited[key] = st.text_area(label, value=data.get(key, ""), height=105, key=f"{idx}-{key}")
                elif key == "customer_po_number":
                    edited[key] = st.text_input(label, value=data.get(key, ""), key=f"{idx}-{key}", help="Use the Fisher Scientific order number. It starts with DR.")
                else:
                    edited[key] = st.text_input(label, value=data.get(key, ""), key=f"{idx}-{key}")
            current_po = str(edited.get("customer_po_number", "")).strip()
            if is_thermo and current_po and not current_po.upper().startswith("DR"):
                st.warning("Customer PO numbers start with DR. Check this against the Fisher Scientific order number.")

        with col2:
            st.markdown("##### Products")
            df = pd.DataFrame(data.get("line_items") or [], columns=["quantity", "item", "unit_price", "amount"])
            edited_df = st.data_editor(df, num_rows="dynamic", width="stretch", key=f"items-{idx}")
            st.caption(f"Text source: {rotation}")
            with st.expander("Extracted text"):
                st.text_area("Extracted text", value=text[:20000], height=320, key=f"raw-{idx}", label_visibility="collapsed")

        line_items = edited_df.fillna("").to_dict(orient="records")
        final_invoice = {**data, **edited, "line_items": line_items}
        final_invoice["issue_date"] = invoice_date_obj.isoformat()
        final_invoice["shipping_date"] = per_shipping_date_obj.isoformat()
        final_invoice["sales_rep"] = sales_rep_default
        final_invoice["order_number"] = per_internal_order_number
        final_invoice["tracking_number"] = per_tracking_number
        context = build_context(sender, final_invoice, per_shipping_cost)

        try:
            with st.spinner(f"Building the invoice for {pdf.name}"):
                docx_bytes = generate_docx(template_bytes, context)
            output_name = f"{safe_filename(final_invoice.get('customer_po_number') or final_invoice.get('order_number') or pdf.name)}_simplex_invoice.docx"
            all_generated[output_name] = docx_bytes
            st.download_button(
                "Download invoice",
                docx_bytes,
                file_name=output_name,
                mime=DOCX_MIME,
                key=f"download-{idx}",
                disabled=not reviewed_ack,
                type="primary",
            )
        except Exception as exc:
            st.error(f"Could not build the invoice for {pdf.name}: {exc}")

    overall_progress.empty()

    if all_generated:
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            for filename, content in all_generated.items():
                zf.writestr(filename, content)
        st.divider()
        st.download_button("Download all invoices (.zip)", zip_buffer.getvalue(), file_name="generated_simplex_invoices.zip", mime="application/zip", disabled=not reviewed_ack, type="primary")


CUSTOM_PRODUCT = "Custom product"
OTHER_SERVICE = "Other service"
NO_SHIPPING_LINE = "No shipping line"
DESTINATIONS = ("United States", "International")


@st.cache_data(ttl=900, show_spinner=False)
def load_catalog() -> tuple[tuple[Product, ...], str]:
    try:
        return fetch_catalog(), f"List prices from simplexsciences.com as of {datetime.now():%H:%M}. Unit prices can be edited."
    except Exception:
        return BUNDLED_CATALOG, (
            f"simplexsciences.com could not be reached. Showing saved list prices from {BUNDLED_CATALOG_DATE}."
        )


def _add_quote_line() -> None:
    ids = st.session_state["quote-line-ids"]
    ids.append(max(ids) + 1)


def _remove_quote_line(line_id: int) -> None:
    ids = st.session_state["quote-line-ids"]
    if line_id in ids and len(ids) > 1:
        ids.remove(line_id)


def _render_quote_lines(catalog: dict[str, Product]) -> list[QuoteLine]:
    ids = st.session_state.setdefault("quote-line-ids", [0])
    choices = list(catalog) + [CUSTOM_PRODUCT]
    lines = []
    for number, line_id in enumerate(ids, start=1):
        with st.container(key=f"quote-line-{line_id}"):
            cols = st.columns([3.4, 1.4, 0.9, 1.1], vertical_alignment="bottom")
            choice = cols[0].selectbox(f"Product {number}", choices, key=f"quote-product-{line_id}")
            product = catalog.get(choice)
            if product is None:
                description = cols[1].text_input("Description", key=f"quote-custom-{line_id}", placeholder="15 nt custom ssDNA marker (100µL)")
                list_price, price_key = 0.0, "custom"
            else:
                option = product.option_labels[0]
                if len(product.options) > 1:
                    option = cols[1].selectbox("Dye", product.option_labels, key=f"quote-option-{line_id}-{choice}")
                description = product.describe(option)
                list_price, price_key = float(product.price_for(option)), f"{choice}-{option}"
            quantity = cols[2].number_input("Quantity", min_value=1, step=1, value=1, key=f"quote-qty-{line_id}")
            unit_price = cols[3].number_input(
                "Unit price", min_value=0.0, step=1.0, format="%.2f", value=list_price,
                key=f"quote-price-{line_id}-{price_key}-{list_price}",
            )
            line = QuoteLine(" ".join(str(description).split()), int(quantity), parse_money(unit_price) or Decimal("0.00"))
            meta = st.container(horizontal=True, vertical_alignment="center", gap="medium")
            note = f"{line.quantity} × {format_money(line.unit_price)} = {format_money(line.amount)}".replace("$", "\\$")
            if line.unit_price == 0:
                note += ". Unit price is \\$0.00."
            meta.caption(note, width="content")
            if product is not None:
                meta.link_button("Product page", product.url, type="tertiary")
            meta.button("Remove", key=f"quote-remove-{line_id}", type="tertiary", on_click=_remove_quote_line, args=(line_id,), disabled=len(ids) == 1)
        lines.append(line)
    st.button("Add product", type="tertiary", on_click=_add_quote_line)
    return lines


def _totals_html(quote: Quote, service: str, shipping_price) -> str:
    if shipping_price is not None:
        shipping = format_money(shipping_price)
    else:
        shipping = "Excluded" if service else "Not on quote"
    total_label = "Total excluding shipping" if quote.excludes_shipping else "Total"
    rows = [("Products", format_money(quote.subtotal)), ("Shipping", shipping)]
    body = "".join(f"<tr><td>{a}</td><td>{b}</td></tr>" for a, b in rows)
    return f'<table class="ss-totals">{body}<tr class="ss-total"><td>{total_label}</td><td>{format_money(quote.total)} USD</td></tr></table>'


def render_quote_tab(sender: dict[str, str], default_sales_rep: str) -> None:
    if st.session_state.get("quote-rep-default") != default_sales_rep:
        st.session_state["quote-rep-default"] = default_sales_rep
        st.session_state["quote-sales-rep"] = default_sales_rep

    st.subheader("Customer", anchor=False)
    c1, c2, c3 = st.columns([1.6, 1, 1])
    issued_to = c1.text_input("Issued to", key="quote-issued-to", placeholder="Company or contact name")
    issue_date = c2.date_input("Issue date", value=date.today(), key="quote-issue-date")
    sales_rep = c3.text_input("Sales rep", key="quote-sales-rep")
    ship_to = st.text_area("Shipping address", key="quote-ship-to", height=130, placeholder="Company\nStreet\nCity, postal code\nCountry")
    destination = st.radio("Destination", DESTINATIONS, horizontal=True, key="quote-destination")

    st.subheader("Products", anchor=False)
    if "quote-catalog" not in st.session_state:
        st.session_state["quote-catalog"] = load_catalog()
    catalog, catalog_note = st.session_state["quote-catalog"]
    n1 = st.container(horizontal=True, vertical_alignment="center", gap="medium")
    n1.caption(catalog_note, width="content")
    if n1.button("Update prices", type="tertiary", key="quote-refresh-prices"):
        load_catalog.clear()
        st.session_state["quote-catalog"] = load_catalog()
        st.rerun()
    lines = _render_quote_lines({product.name: product for product in catalog})

    st.subheader("Shipping", anchor=False)
    default_service = DOMESTIC_SERVICE if destination == DESTINATIONS[0] else INTERNATIONAL_SERVICE
    services = list(SHIPPING_SERVICES) + [OTHER_SERVICE, NO_SHIPPING_LINE]
    s1, s2 = st.columns([1.4, 1])
    service_choice = s1.selectbox("FedEx service", services, index=services.index(default_service), key=f"quote-service-{destination}")
    service = service_choice
    if service == OTHER_SERVICE:
        service = s1.text_input("Service name on the quote", key="quote-service-other").strip()
    elif service == NO_SHIPPING_LINE:
        service = ""
    price_text = s2.text_input("Shipping price (USD)", key="quote-shipping-price", placeholder="32.38", disabled=not service)
    ship_from = ", ".join(line.strip() for line in sender.get("company_address", "").splitlines() if line.strip())
    st.caption(
        f"Rate a package from {ship_from} to the shipping address in the "
        f"[FedEx rate calculator]({FEDEX_RATES_URL}), then enter the price for the selected service. "
        "Leave the price blank to quote without shipping."
    )

    problems = []
    shipping_price = None
    if service:
        try:
            shipping_price = parse_money(price_text)
        except ValueError:
            problems.append("Enter the shipping price as a number, for example 32.38.")

    quote = Quote(
        issued_to=issued_to.strip(),
        issue_date=issue_date.isoformat(),
        sales_rep=sales_rep.strip(),
        ship_to=ship_to,
        lines=lines,
        shipping_service=service,
        shipping_price=shipping_price,
    )
    if service_choice == OTHER_SERVICE and not service:
        problems.append("Enter the shipping service name, or choose No shipping line.")
    problems = validate_quote(quote) + problems

    st.subheader("Review", anchor=False)
    st.html(_totals_html(quote, service, shipping_price))

    default_label = default_file_label(quote.issued_to, quote.ship_to)
    if not st.session_state.get("quote-file-label-edited") and st.session_state.get("quote-file-label") != default_label:
        st.session_state["quote-file-label"] = default_label
    label = st.text_input(
        "Customer name in file name", key="quote-file-label", width=420,
        on_change=lambda: st.session_state.update({"quote-file-label-edited": True}),
    )
    file_name = quote_filename(quote.issue_date, label)

    with st.expander("Quote template"):
        uploaded_quote_template = st.file_uploader("Quote template (.docx)", type=["docx"], key="quote-template")
        st.caption("Optional. Without an upload, the app uses simplex_quote_template.docx. A previously sent quote also works.")
    template_bytes = uploaded_quote_template.getvalue() if uploaded_quote_template else DEFAULT_QUOTE_TEMPLATE.read_bytes()

    docx_bytes = None
    if problems:
        st.info("Still needed:\n" + "\n".join(f"- {problem}" for problem in problems))
    else:
        try:
            docx_bytes = generate_quote_docx(template_bytes, quote, sender)
        except QuoteTemplateError as exc:
            st.error(str(exc))
        except Exception as exc:
            st.error(f"Could not build the quote: {exc}")
    st.download_button(
        "Download quote", docx_bytes or b"", file_name=file_name, mime=DOCX_MIME,
        disabled=docx_bytes is None, key="quote-download", type="primary",
    )
    st.caption(f"Saves as {file_name}")


if __name__ == "__main__":
    main()
